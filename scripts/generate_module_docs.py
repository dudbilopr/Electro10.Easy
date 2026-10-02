#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
scripts/generate_module_docs.py
Generador automatico de documentacion Word (.docx) para Electro10.Easy
Modulos 0-13 — Curso de Electromagnetismo UNAB
"""
import os, sys
from pathlib import Path
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.exit("ERROR: python-docx no instalado.")

try:
    import latex2mathml.converter as l2m
    HAS_L2M = True
except ImportError:
    HAS_L2M = False

try:
    from lxml import etree
except ImportError:
    sys.exit("ERROR: lxml no instalado.")

UNAB_BLUE    = RGBColor(0x00, 0x38, 0x65)
UNAB_ORANGE  = RGBColor(0xE8, 0x5C, 0x0B)
DARK_GRAY    = RGBColor(0x33, 0x33, 0x33)
MID_GRAY     = RGBColor(0x66, 0x66, 0x66)

MODULES_DATA = [
    {"id":0,"titulo":"Modulo 0: Examen Diagnostico de Presaberes",
     "descripcion":"Evaluacion diagnostica holistica de estilos de aprendizaje (CHAEA/Kolb) y competencias matematicas previas.",
     "subtemas":["Diagnostico de estilos de aprendizaje: CHAEA y modelo de Kolb",
                 "Evaluacion diagnostica de competencias matematicas: algebra vectorial y calculo",
                 "Rutas de nivelacion personalizada"],"formulas":[]},
    {"id":1,"titulo":"Modulo 1: Ley de Coulomb",
     "descripcion":"Cuantizacion de la carga electrica, polarizacion y fuerza electrostatica vectorial.",
     "subtemas":["Naturaleza atomica de la carga electrica y principio de conservacion",
                 "Cuantizacion de la carga y el experimento de Millikan",
                 "Conductores, aislantes y polarizacion dielectrica",
                 "Formulacion escalar de la Ley de Coulomb",
                 "Formulacion vectorial rigurosa de la Ley de Coulomb",
                 "Principio de superposicion lineal para distribuciones discretas",
                 "Equilibrio electrostatico en configuraciones de cargas puntuales"],
     "formulas":[
         {"nombre":"Fuerza de Coulomb (escalar)","latex":"F = k_e |q_1 q_2| / r^2",
          "descripcion":"Magnitud de la fuerza entre dos cargas puntuales."},
         {"nombre":"Ley de Coulomb (vectorial)","latex":"F_12 = k_e * q1*q2/r^2 * r_hat_12",
          "descripcion":"Fuerza vectorial que ejerce q1 sobre q2."},
         {"nombre":"Constante de Coulomb","latex":"k_e = 1/(4*pi*epsilon_0) ≈ 8.988e9 N·m²/C²",
          "descripcion":"Constante de proporcionalidad en unidades SI."},
         {"nombre":"Superposicion","latex":"F_neta = Sum_i F_i",
          "descripcion":"Fuerza total como suma vectorial de contribuciones individuales."}]},
    {"id":2,"titulo":"Modulo 2: Campo Electrico",
     "descripcion":"Campo electrico vectorial de distribuciones discretas y continuas de carga.",
     "subtemas":["Definicion de campo electrico E = F/q0","Campo de carga puntual",
                 "Superposicion vectorial","Dipolo electrico",
                 "Distribucion discreta de cargas","Lineas de campo",
                 "Conductor en equilibrio"],
     "formulas":[
         {"nombre":"Campo electrico","latex":"E = k_e * q/r^2 * r_hat",
          "descripcion":"Campo como fuerza por unidad de carga de prueba."},
         {"nombre":"Superposicion de E","latex":"E_total = Sum_i k_e*qi/ri^2 * r_hat_i",
          "descripcion":"Campo resultante de multiples cargas."},
         {"nombre":"Momento dipolar","latex":"p = q*d","descripcion":"Vector dipolo electrico."},
         {"nombre":"Torque sobre dipolo","latex":"tau = p x E","descripcion":"Torque en campo uniforme."}]},
    {"id":3,"titulo":"Modulo 3: Potencial Electrico",
     "descripcion":"Energia potencial electrostatica, potencial V y superficies equipotenciales.",
     "subtemas":["Trabajo de la fuerza electrica","Potencial electrico V",
                 "Relacion E = -grad(V)","Superficies equipotenciales",
                 "Potencial de distribuciones discretas","Energia del sistema",
                 "Potencial en conductores"],
     "formulas":[
         {"nombre":"Potencial de carga puntual","latex":"V = k_e * q/r",
          "descripcion":"Potencial a distancia r de carga q."},
         {"nombre":"Trabajo y potencial","latex":"W = q0*(VA - VB)",
          "descripcion":"Trabajo al mover carga q0 del punto A al B."},
         {"nombre":"Campo desde potencial","latex":"E = -grad(V)",
          "descripcion":"Campo electrico como gradiente negativo del potencial."},
         {"nombre":"Energia del sistema","latex":"U = k_e * Sum_ij qi*qj/rij",
          "descripcion":"Energia total de configuracion de N cargas."}]},
    {"id":4,"titulo":"Modulo 4: Distribucion de Cargas Continuas",
     "descripcion":"Campo electrico por integracion de densidades lineales, superficiales y volumetricas.",
     "subtemas":["Densidades: lambda, sigma, rho","Campo de varilla cargada",
                 "Campo de anillo circular","Campo de disco",
                 "Campo de lamina infinita","Placas paralelas","Integracion analitica"],
     "formulas":[
         {"nombre":"Elemento diferencial de campo","latex":"dE = k_e * dq/r^2 * r_hat",
          "descripcion":"Contribucion diferencial de elemento dq."},
         {"nombre":"Densidades de carga","latex":"dq = lambda*dl = sigma*dA = rho*dV",
          "descripcion":"Elemento de carga segun dimension."},
         {"nombre":"Campo de hilo infinito","latex":"E = lambda/(2*pi*epsilon_0*r)",
          "descripcion":"Campo radial a distancia r de hilo infinito."},
         {"nombre":"Campo de placa infinita","latex":"E = sigma/(2*epsilon_0)",
          "descripcion":"Campo uniforme de lamina infinita."}]},
    {"id":5,"titulo":"Modulo 5: Ley de Gauss",
     "descripcion":"Flujo del campo electrico y calculo de campo mediante simetrias gaussianas.",
     "subtemas":["Flujo electrico por superficies","Vector diferencial dA",
                 "Ley de Gauss integral","Simetria esferica","Simetria cilindrica",
                 "Simetria planar","Conductores y Jaula de Faraday"],
     "formulas":[
         {"nombre":"Flujo electrico","latex":"Phi_E = integral E·dA",
          "descripcion":"Integral del campo a traves de superficie S."},
         {"nombre":"Ley de Gauss","latex":"closed_integral E·dA = Q_enc/epsilon_0",
          "descripcion":"Flujo total proporcional a carga encerrada."},
         {"nombre":"Campo esfera (r>R)","latex":"E = Q/(4*pi*epsilon_0*r^2)",
          "descripcion":"Campo exterior de esfera simetrica."},
         {"nombre":"Campo esfera (r<R)","latex":"E = rho*r/(3*epsilon_0)",
          "descripcion":"Campo interior de esfera no conductora uniforme."}]},
    {"id":6,"titulo":"Modulo 6: Capacitancia y Dielectricos",
     "descripcion":"Almacenamiento de carga y energia, polarizacion en dielectricos y redes capacitivas.",
     "subtemas":["Capacitancia C = Q/V","Placas paralelas C = epsilon_0*A/d",
                 "Capacitores cilindricos y esfericos","Arreglos serie y paralelo",
                 "Energia U = 1/2*CV^2","Dielectricos y constante kappa","Ruptura dielectrica"],
     "formulas":[
         {"nombre":"Definicion de capacitancia","latex":"C = Q/V",
          "descripcion":"Capacitancia como razon carga/potencial."},
         {"nombre":"Placas paralelas","latex":"C = epsilon_0*A/d",
          "descripcion":"Capacitancia de placas plano-paralelas en vacio."},
         {"nombre":"Capacitores en paralelo","latex":"C_eq = Sum_i C_i",
          "descripcion":"Capacitancia equivalente en paralelo."},
         {"nombre":"Capacitores en serie","latex":"1/C_eq = Sum_i 1/C_i",
          "descripcion":"Capacitancia equivalente en serie."},
         {"nombre":"Energia almacenada","latex":"U = (1/2)*C*V^2 = Q^2/(2C)",
          "descripcion":"Energia potencial en capacitor cargado."},
         {"nombre":"Con dielectrico","latex":"C = kappa*epsilon_0*A/d",
          "descripcion":"Efecto del dielectrico con constante kappa."}]},
    {"id":7,"titulo":"Modulo 7: Circuitos de Corriente Continua (CC)",
     "descripcion":"Analisis de redes resistivas mediante leyes de Kirchhoff.",
     "subtemas":["Corriente I = dQ/dt","Ley de Ohm V = IR",
                 "Potencia P = IV","Serie y paralelo","Leyes de Kirchhoff LCK y LVK",
                 "Metodo de mallas y nodos","Teoremas de Thevenin y Norton"],
     "formulas":[
         {"nombre":"Corriente electrica","latex":"I = dQ/dt",
          "descripcion":"Tasa de flujo de carga electrica."},
         {"nombre":"Ley de Ohm","latex":"V = I*R",
          "descripcion":"Diferencia de potencial proporcional a corriente."},
         {"nombre":"Resistores en serie","latex":"R_eq = Sum_i R_i",
          "descripcion":"Resistencia equivalente en serie."},
         {"nombre":"Resistores en paralelo","latex":"1/R_eq = Sum_i 1/R_i",
          "descripcion":"Resistencia equivalente en paralelo."},
         {"nombre":"Ley de Kirchhoff de corrientes","latex":"Sum I_k = 0",
          "descripcion":"Suma de corrientes en nodo es cero."},
         {"nombre":"Ley de Kirchhoff de voltajes","latex":"Sum V_k = 0",
          "descripcion":"Suma de voltajes en lazo cerrado es cero."},
         {"nombre":"Potencia electrica","latex":"P = I*V = I^2*R = V^2/R",
          "descripcion":"Potencia disipada por elemento del circuito."}]},
    {"id":8,"titulo":"Modulo 8: Ley de Lorentz y Campo Magnetico",
     "descripcion":"Fuerza magnetica sobre cargas y conductores, movimiento en campos cruzados.",
     "subtemas":["Campo magnetico B","Fuerza de Lorentz F = q(E + vxB)",
                 "Movimiento circular","Espectrometro de masas","Fuerza sobre conductores",
                 "Momento dipolar magnetico","Efecto Hall"],
     "formulas":[
         {"nombre":"Fuerza de Lorentz","latex":"F = q*(E + v x B)",
          "descripcion":"Fuerza total sobre carga en campos cruzados."},
         {"nombre":"Fuerza magnetica","latex":"F_B = q*v x B",
          "descripcion":"Componente magnetica de la fuerza de Lorentz."},
         {"nombre":"Radio de orbita","latex":"r = m*v/(|q|*B)",
          "descripcion":"Radio de trayectoria circular en campo B uniforme."},
         {"nombre":"Fuerza sobre conductor","latex":"dF = I*dl x B",
          "descripcion":"Fuerza diferencial sobre segmento con corriente."}]},
    {"id":9,"titulo":"Modulo 9: Ley de Biot-Savart",
     "descripcion":"Campo magnetico generado por corrientes estacionarias.",
     "subtemas":["Biot-Savart diferencial","Campo de conductor recto",
                 "Campo de espira circular","Solenoides y toroides",
                 "Fuerza entre conductores","Momento dipolar magnetico","Definicion del Ampere"],
     "formulas":[
         {"nombre":"Ley de Biot-Savart","latex":"dB = (mu_0/4pi) * I*dl x r_hat / r^2",
          "descripcion":"Campo diferencial de elemento de corriente I*dl."},
         {"nombre":"Campo de hilo infinito","latex":"B = mu_0*I/(2*pi*r)",
          "descripcion":"Campo a distancia r de conductor recto infinito."},
         {"nombre":"Campo en centro de espira","latex":"B = mu_0*I/(2*R)",
          "descripcion":"Campo magnetico en el centro de espira de radio R."},
         {"nombre":"Campo axial de espira","latex":"B_x = mu_0*I*R^2 / (2*(R^2+x^2)^(3/2))",
          "descripcion":"Campo en el eje de una espira a distancia x del centro."}]},
    {"id":10,"titulo":"Modulo 10: Ley de Ampere",
     "descripcion":"Circulacion del campo magnetico y corriente de desplazamiento de Maxwell.",
     "subtemas":["Ley de Ampere integral","Caminos amperianos","Solenoides",
                 "Toroides","Conductores coaxiales","Corriente de desplazamiento",
                 "Ampere-Maxwell generalizada"],
     "formulas":[
         {"nombre":"Ley de Ampere","latex":"closed_integral B·dl = mu_0*I_enc",
          "descripcion":"Circulacion de B proporcional a corriente encerrada."},
         {"nombre":"Campo en solenoide ideal","latex":"B = mu_0*n*I",
          "descripcion":"Campo uniforme en solenoide con n espiras/m."},
         {"nombre":"Corriente de desplazamiento","latex":"I_D = epsilon_0 * d(Phi_E)/dt",
          "descripcion":"Corriente de desplazamiento de Maxwell."},
         {"nombre":"Ampere-Maxwell","latex":"closed_integral B·dl = mu_0*(I + epsilon_0*dPhi_E/dt)",
          "descripcion":"Generalizacion incluyendo corriente de desplazamiento."}]},
    {"id":11,"titulo":"Modulo 11: Ley de Faraday",
     "descripcion":"FEM inducida por flujo magnetico variable, ley de Lenz e inductancia.",
     "subtemas":["Flujo magnetico Phi_B","FEM = -dPhi_B/dt","Ley de Lenz",
                 "FEM de movimiento","Inductancia mutua M","Autoinductancia L",
                 "Transformadores ideales"],
     "formulas":[
         {"nombre":"Flujo magnetico","latex":"Phi_B = integral B·dA",
          "descripcion":"Integral de superficie del campo magnetico."},
         {"nombre":"Ley de Faraday","latex":"epsilon = -N*d(Phi_B)/dt",
          "descripcion":"FEM inducida en N espiras por flujo variable."},
         {"nombre":"FEM de movimiento","latex":"epsilon = B*L*v",
          "descripcion":"FEM en conductor de longitud L a velocidad v en campo B."},
         {"nombre":"Energia en inductor","latex":"U_B = (1/2)*L*I^2",
          "descripcion":"Energia magnetica almacenada en inductor L."}]},
    {"id":12,"titulo":"Modulo 12: Ecuaciones de Maxwell y Ondas EM",
     "descripcion":"Sintesis del electromagnetismo clasico y ondas electromagneticas.",
     "subtemas":["Maxwell integral","Maxwell diferencial","Corriente de desplazamiento",
                 "Ecuacion de onda","Velocidad c = 1/sqrt(epsilon_0*mu_0)",
                 "Vector de Poynting","Espectro EM y polarizacion"],
     "formulas":[
         {"nombre":"Maxwell I — Gauss electrica","latex":"closed_integral E·dA = Q_enc/epsilon_0",
          "descripcion":"Flujo electrico total proporcional a carga encerrada."},
         {"nombre":"Maxwell II — Gauss magnetica","latex":"closed_integral B·dA = 0",
          "descripcion":"No existen monopolos magneticos."},
         {"nombre":"Maxwell III — Faraday","latex":"closed_integral E·dl = -d(Phi_B)/dt",
          "descripcion":"Campo B variable induce campo E rotacional."},
         {"nombre":"Maxwell IV — Ampere-Maxwell","latex":"closed_integral B·dl = mu_0*I + mu_0*epsilon_0*d(Phi_E)/dt",
          "descripcion":"Corrientes reales y de desplazamiento generan campo B."},
         {"nombre":"Velocidad de la luz","latex":"c = 1/sqrt(epsilon_0*mu_0) ≈ 3e8 m/s",
          "descripcion":"Derivada de Maxwell: velocidad de EM en vacio."},
         {"nombre":"Vector de Poynting","latex":"S = (1/mu_0) * E x B",
          "descripcion":"Densidad de flujo de energia EM [W/m^2]."}]},
    {"id":13,"titulo":"Modulo 13: Circuitos RL, LC, RLC y CA",
     "descripcion":"Respuesta transitoria, oscilaciones, fasores y resonancia electrica.",
     "subtemas":["Circuito RL: tau = L/R","Circuito LC: oscilaciones",
                 "Circuito RLC amortiguado","Fasores","Reactancia XL y XC, impedancia Z",
                 "Resonancia y factor Q","Potencia activa P y factor de potencia"],
     "formulas":[
         {"nombre":"Frecuencia de resonancia","latex":"omega_0 = 1/sqrt(L*C)",
          "descripcion":"Frecuencia natural de oscilacion del circuito LC."},
         {"nombre":"Impedancia compleja","latex":"Z = R + j*(omega*L - 1/(omega*C))",
          "descripcion":"Impedancia total del circuito RLC en regimen senoidal."},
         {"nombre":"Ley de Ohm fasorial","latex":"V_tilde = I_tilde * Z",
          "descripcion":"Version fasorial de la ley de Ohm en CA."},
         {"nombre":"Potencia activa","latex":"P = Vrms*Irms*cos(phi)",
          "descripcion":"Potencia media disipada en el resistor."},
         {"nombre":"Factor de calidad","latex":"Q = omega_0*L/R = (1/R)*sqrt(L/C)",
          "descripcion":"Selectividad del circuito resonante."}]}
]

def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1");    bottom.set(qn("w:color"), "003865")
    pBdr.append(bottom); pPr.append(pBdr)

def add_math_text(p, latex_text):
    run = p.add_run(latex_text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(11)
    run.font.italic = True

def generate_doc(mod, out_dir):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Cm(2.5); sec.bottom_margin = Cm(2.5)
        sec.left_margin   = Cm(3.0); sec.right_margin  = Cm(2.5)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    # ── Portada ──
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("UNIVERSIDAD AUTONOMA DE BUCARAMANGA")
    r.bold = True; r.font.size = Pt(14); r.font.color.rgb = UNAB_BLUE

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Introduccion al Electromagnetismo — CDAT")
    r.font.size = Pt(11); r.font.color.rgb = MID_GRAY

    doc.add_paragraph(); add_hr(doc); doc.add_paragraph()

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(mod["titulo"])
    r.bold = True; r.font.size = Pt(20); r.font.color.rgb = UNAB_BLUE

    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(mod["descripcion"])
    r.font.size = Pt(12); r.font.color.rgb = MID_GRAY; r.font.italic = True

    doc.add_paragraph(); add_hr(doc); doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Generado: {datetime.now().strftime('%d/%m/%Y')}")
    r.font.size = Pt(9); r.font.color.rgb = MID_GRAY
    doc.add_page_break()

    # ── Seccion 1: Subtemas ──
    p = doc.add_paragraph()
    r = p.add_run("1. Competencias y Contenidos Tematicos")
    r.bold = True; r.font.size = Pt(14); r.font.color.rgb = UNAB_BLUE
    doc.add_paragraph()
    for idx, st in enumerate(mod["subtemas"], 1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(st); r.font.size = Pt(11); r.font.color.rgb = DARK_GRAY

    # ── Seccion 2: Formulario ──
    if mod["formulas"]:
        doc.add_paragraph(); add_hr(doc); doc.add_paragraph()
        p = doc.add_paragraph()
        r = p.add_run("2. Formulario Oficial del Modulo")
        r.bold = True; r.font.size = Pt(14); r.font.color.rgb = UNAB_BLUE
        doc.add_paragraph()
        for idx, f in enumerate(mod["formulas"], 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            r1 = p.add_run(f"{idx}. "); r1.bold = True; r1.font.color.rgb = UNAB_ORANGE; r1.font.size = Pt(11)
            r2 = p.add_run(f["nombre"]); r2.bold = True; r2.font.size = Pt(11); r2.font.color.rgb = UNAB_BLUE

            p_eq = doc.add_paragraph(); p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_eq.paragraph_format.space_before = Pt(2); p_eq.paragraph_format.space_after = Pt(2)
            add_math_text(p_eq, f["latex"])

            p_d = doc.add_paragraph(f"   {f['descripcion']}")
            p_d.paragraph_format.space_after = Pt(8)
            for r in p_d.runs:
                r.font.color.rgb = MID_GRAY; r.font.size = Pt(10); r.font.italic = True

    # ── Seccion 3: Estructura 9 bloques ──
    doc.add_paragraph(); add_hr(doc); doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("3. Estructura Pedagogica — 9 Bloques de Aprendizaje")
    r.bold = True; r.font.size = Pt(14); r.font.color.rgb = UNAB_BLUE
    doc.add_paragraph()

    bloques = [
        ("Bloque 1", "Videos Magistrales",          "Fundamentacion teorica y deduccion matematica."),
        ("Bloque 2", "Diapositivas de Apoyo",        "Material visual interactivo de la clase magistral."),
        ("Bloque 3", "Laboratorio Virtual",          "Simuladores interactivos nativos de la plataforma."),
        ("Bloque 4", "Physics Quest",                "Juego educativo de misiones fisicas (5 retos)."),
        ("Bloque 5", "Taller Practico",              "Problemas de nivel universitario Sears-Zemansky."),
        ("Bloque 6", "Quiz Adaptativo",             "Evaluacion formativa en 3 niveles cognitivos."),
        ("Bloque 7", "Evaluacion Sumativa",          "Cuestionario oficial de acreditacion del modulo."),
        ("Bloque 8", "Asistente IA (NotebookLM)",   "IA curada para consultas analiticas y derivaciones."),
        ("Bloque 9", "Repositorio Documental",       "Referencias bibliograficas y simuladores externos."),
    ]
    for code, name, desc in bloques:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.5)
        r1 = p.add_run(f"{code}: "); r1.bold = True; r1.font.color.rgb = UNAB_ORANGE; r1.font.size = Pt(11)
        r2 = p.add_run(f"{name} — "); r2.bold = True; r2.font.size = Pt(11); r2.font.color.rgb = DARK_GRAY
        r3 = p.add_run(desc); r3.font.size = Pt(11); r3.font.color.rgb = MID_GRAY

    # ── Guardar ──
    fname = f"Modulo_{mod['id']:02d}_{mod['titulo'].split(':')[0].strip().replace(' ','_')}.docx"
    out_path = out_dir / fname
    doc.save(str(out_path))
    return out_path

def main():
    repo = Path(__file__).parent.parent
    out  = repo / "documentacion"
    out.mkdir(exist_ok=True)
    print(f"\n{'='*55}\n  Electro10.Easy — Word Doc Generator\n  Destino: {out}\n{'='*55}\n")
    ok, err = [], []
    for mod in MODULES_DATA:
        try:
            p = generate_doc(mod, out)
            print(f"  [OK] M{mod['id']:02d} -> {p.name}")
            ok.append(p)
        except Exception as e:
            print(f"  [ERROR] M{mod['id']:02d} -> {e}")
            err.append((mod['id'], str(e)))
    print(f"\n{'='*55}\n  Exitosos: {len(ok)}/{len(MODULES_DATA)}")
    if err:
        for eid, em in err: print(f"  ERROR M{eid:02d}: {em}")
    print(f"{'='*55}\n")

if __name__ == "__main__":
    main()
